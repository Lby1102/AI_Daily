# -*- coding: utf-8 -*-
"""从（已增强的）article.html 生成 docx 新闻稿，供宣传部审核。

HTML 结构来自模块 2~4 固定模板：h1/.introduction/.feature-article(h2/
.article-lead/.article-body p/.source)/.conclusion，图片用 <img> 本地相对路径。
"""
from __future__ import annotations

from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Cm, Pt, RGBColor


def _resolve_img(path: str, base_dir: Path) -> Path | None:
    p = Path(path)
    if not p.is_absolute():
        p = base_dir / p
    return p if p.exists() else None


def build_docx(article_html: Path, docx_path: Path, base_dir: Path | None = None) -> Path:
    """article_html 所在目录作为相对图片基准；默认即其父目录。"""
    base_dir = base_dir or article_html.parent
    soup = BeautifulSoup(article_html.read_text(encoding="utf-8"), "html.parser")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    try:  # 中文字体映射
        from docx.oxml.ns import qn
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    except Exception:
        pass
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    h1 = soup.find("h1")
    doc.add_heading((h1.get_text(" ", strip=True) if h1 else "AI 资讯"), level=0)

    meta_p = soup.select_one("p.meta")
    if meta_p:
        p = doc.add_paragraph(meta_p.get_text(" ", strip=True))
        p.runs[0].font.color.rgb = RGBColor(0x68, 0x73, 0x86)
        p.runs[0].font.size = Pt(9)

    intro = soup.select_one(".introduction")
    if intro:
        doc.add_paragraph(intro.get_text(" ", strip=True))

    for article in soup.select("article.feature-article"):
        h2 = article.find("h2")
        if h2:
            doc.add_heading(h2.get_text(" ", strip=True), level=1)
        lead = article.select_one(".article-lead")
        if lead:
            pl = doc.add_paragraph(lead.get_text(" ", strip=True))
            for run in pl.runs:
                run.font.bold = True
        for para in article.select(".article-body > p"):
            doc.add_paragraph(para.get_text(" ", strip=True))
        for img in article.find_all("img"):
            src = img.get("src", "")
            if not src or src.startswith(("http://", "https://", "data:")):
                continue
            local = _resolve_img(src, base_dir)
            if not local:
                continue
            try:
                doc.add_picture(str(local), width=Cm(15.0))
                doc.paragraphs[-1].alignment = 1  # 居中
            except Exception:
                continue
        source = article.select_one(".source")
        if source:
            doc.add_paragraph(source.get_text(" ", strip=True))

    conclusion = soup.select_one(".conclusion")
    if conclusion:
        h = conclusion.find("h2")
        if h:
            doc.add_heading(h.get_text(" ", strip=True), level=1)
        for p in conclusion.find_all("p"):
            doc.add_paragraph(p.get_text(" ", strip=True))

    footer = soup.find("footer")
    if footer:
        doc.add_paragraph(footer.get_text(" ", strip=True))

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    return docx_path
